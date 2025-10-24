import os
import pytesseract
from pdf2image import convert_from_path
from openpyxl import Workbook, load_workbook

# Path ng folders
pdf_folder = r"C:\Users\ASTA\Desktop\OT_NA_DAPAT_BAYARAN\OVERTIME REQUEST FORM"
summary_file = r"C:\Users\ASTA\Desktop\OT_NA_DAPAT_BAYARAN\OT_SUMMARY.xlsx"

# OCR config (ensure naka-install ang Tesseract)
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

# Kung wala pang summary file, create one
if not os.path.exists(summary_file):
    wb = Workbook()
    ws = wb.active
    ws.append(["File", "Employee Name", "Date", "OT Hours"])  # header
    wb.save(summary_file)

wb = load_workbook(summary_file)
ws = wb.active

# Loop sa lahat ng PDF
for file in os.listdir(pdf_folder):
    if file.endswith(".pdf"):
        pdf_path = os.path.join(pdf_folder, file)
        images = convert_from_path(pdf_path)

        text = ""
        for img in images:
            text += pytesseract.image_to_string(img)

        # --- Example Parsing (edit mo depende sa format ng form) ---
        employee = "Not Found"
        date = "Not Found"
        hours = "0"

        for line in text.splitlines():
            if "Employee" in line:
                employee = line.split(":")[-1].strip()
            if "Date" in line:
                date = line.split(":")[-1].strip()
            if "Hours" in line:
                hours = line.split(":")[-1].strip()

        # Append sa Excel
        ws.append([file, employee, date, hours])

wb.save(summary_file)
print("OT Summary updated successfully!")
